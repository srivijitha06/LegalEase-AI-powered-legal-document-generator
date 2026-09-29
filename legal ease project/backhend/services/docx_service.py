from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt


def create_docx(
    text,
    document_type,
):

    document = Document()

    section = document.sections[0]

    section.top_margin = Pt(50)
    section.bottom_margin = Pt(50)
    section.left_margin = Pt(60)
    section.right_margin = Pt(60)

    normal = document.styles["Normal"]

    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    title = document.add_paragraph()

    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run(
        "LEGALEASE"
    )

    run.bold = True
    run.font.size = Pt(22)

    subtitle = document.add_paragraph()

    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run(
        document_type.upper()
    )

    run.bold = True
    run.font.size = Pt(14)

    for line in text.splitlines():

        line = line.strip()

        if not line:

            document.add_paragraph()

            continue

        paragraph = document.add_paragraph()

        if (
            line.isupper()
            or (
                len(line) > 2
                and line[0].isdigit()
                and "." in line[:4]
            )
        ):

            run = paragraph.add_run(
                line
            )

            run.bold = True

        else:

            paragraph.add_run(
                line
            )

    footer = section.footer.paragraphs[0]

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    footer_run = footer.add_run(
        "LegalEase | AI-assisted draft"
    )

    footer_run.font.size = Pt(8)

    output = BytesIO()

    document.save(output)

    return output.getvalue()